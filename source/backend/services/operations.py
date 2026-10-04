"""Offline data exchange and transactional SQLite backup/restore."""
import json
import sqlite3
import tempfile
import threading
import uuid
from pathlib import Path
from io import BytesIO
from contextlib import closing
from django.conf import settings
from django.db import connections, transaction
from django.http import FileResponse
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from core import models as m
from core.registry import RESOURCES, READONLY, serializer_for

DATABASE_LOCK=threading.RLock()


class DatabaseGuard:
    """Single-user portable server: serialize requests against restore/backup maintenance."""
    def __init__(self,get_response): self.get_response=get_response
    def __call__(self,request):
        with DATABASE_LOCK: return self.get_response(request)


def audit(user,action,entity='',pk='',details=None):
    m.AuditLog.objects.create(user=user,action=action,entity_type=entity,entity_id=str(pk),details=details or {})


def workbook_response(name,rows,template=False):
    from openpyxl import Workbook
    cls=serializer_for(name)
    fields=[key for key,field in cls().fields.items() if key!='password' and not (template and field.read_only)]
    book=Workbook()
    sheet=book.active
    sheet.title='data'
    sheet.append(fields)
    for row in rows:
        values=[]
        for key in fields:
            value=row.get(key)
            if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
            if isinstance(value,str) and value.startswith(('=','+','-','@')): value="'"+value
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=min(45,max(15,len(str(column[0].value))+3))
    stream=BytesIO()
    book.save(stream)
    stream.seek(0)
    return FileResponse(stream,as_attachment=True,filename=f'{name}_{"template" if template else "export"}.xlsx')


def import_workbook(user,name,upload):
    from openpyxl import load_workbook
    from rest_framework import serializers
    if name not in RESOURCES or name in READONLY or name in ['users','applications','service-requests']:
        raise ValidationError('此资源不支持批量导入。')
    if not upload or upload.size>10*1024*1024 or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('请上传10MB以内的xlsx文件。')
    try:
        import zipfile
        with zipfile.ZipFile(upload) as archive:
            if sum(info.file_size for info in archive.infolist())>50*1024*1024:
                raise ValueError('uncompressed size')
        upload.seek(0)
        book=load_workbook(upload,read_only=True,data_only=False)
        sheet=book.active
        if sheet.max_row>10001 or sheet.max_column>100: raise ValueError('size')
        iterator=sheet.iter_rows(values_only=True)
        headers=list(next(iterator))
    except Exception: raise ValidationError('无法读取工作簿，或超过10000行/100列限制。')
    cls=serializer_for(name)
    allowed=cls().fields
    if len(set(headers))!=len(headers) or any(key not in allowed for key in headers):
        book.close()
        raise ValidationError('表头不匹配，请使用本版本下载的导入模板。')
    batch=m.ImportBatch.objects.create(user=user,import_type=name,file_name=Path(upload.name).name)
    for line,values in enumerate(iterator,2):
        if not any(value is not None for value in values): continue
        batch.total_rows+=1
        try:
            data={key:value for key,value in zip(headers,values) if value is not None}
            for key,field in allowed.items():
                if key in data and isinstance(field,serializers.JSONField) and isinstance(data[key],str): data[key]=json.loads(data[key])
            if any(isinstance(v,str) and v.startswith('=') for v in data.values()): raise ValueError('不允许公式单元格')
            pk=data.pop('id',None)
            obj=RESOURCES[name][0].objects.filter(pk=pk).first() if pk else None
            if pk and not obj: raise ValueError('更新目标ID不存在')
            with transaction.atomic():
                serializer=cls(obj,data=data,partial=obj is not None)
                serializer.is_valid(raise_exception=True)
                saved=serializer.save()
                if isinstance(saved,m.PaperQuestion):
                    from django.db.models import Sum
                    paper=saved.paper
                    paper.total_questions=paper.paper_questions.count()
                    paper.total_score=paper.paper_questions.aggregate(value=Sum('points'))['value'] or 0
                    paper.save(update_fields=['total_questions','total_score'])
            batch.success_count+=1
            batch.duplicate_count+=int(obj is not None)
        except Exception as error:
            batch.failure_count+=1
            batch.error_report.append({'row':line,'error':str(error)[:600]})
    book.close()
    batch.status='completed' if not batch.failure_count else 'partial' if batch.success_count else 'failed'
    batch.finished_at=timezone.now()
    batch.save()
    audit(user,'批量导入',name,batch.pk,{'success':batch.success_count,'failure':batch.failure_count})
    return batch


def backup(user):
    directory=Path(settings.DATA_DIR)/'backups'
    directory.mkdir(exist_ok=True)
    filename=f'backup_{timezone.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}.sqlite3'
    target=directory/filename
    connections.close_all()
    with closing(sqlite3.connect(str(settings.DATABASES['default']['NAME']))) as source,closing(sqlite3.connect(target)) as destination:
        source.backup(destination)
    audit(user,'创建备份','database',filename)
    return filename


def restore(user,upload):
    if not upload or upload.size>100*1024*1024: raise ValidationError('请选择100MB以内的数据库备份。')
    with tempfile.TemporaryDirectory(prefix='exam_restore_') as temporary:
        candidate=Path(temporary)/'candidate.sqlite3'
        with candidate.open('wb') as file:
            for chunk in upload.chunks(): file.write(chunk)
        try:
            with closing(sqlite3.connect(f'{candidate.as_uri()}?mode=ro',uri=True)) as source:
                if source.execute('PRAGMA integrity_check').fetchone()[0]!='ok': raise ValueError('损坏')
                from django.apps import apps
                for model in apps.get_models(include_auto_created=True):
                    columns={row[1] for row in source.execute(f'PRAGMA table_info("{model._meta.db_table}")')}
                    if not {field.column for field in model._meta.local_fields}.issubset(columns): raise ValueError('备份版本不兼容')
                if source.execute('PRAGMA foreign_key_check').fetchone(): raise ValueError('存在损坏的关联数据')
                if not source.execute("SELECT COUNT(*) FROM core_user WHERE role='admin' AND status='active'").fetchone()[0]: raise ValueError('备份没有可用管理员')
                recovery=backup(user)
                connections.close_all()
                with closing(sqlite3.connect(str(settings.DATABASES['default']['NAME']))) as destination: source.backup(destination)
        except (sqlite3.DatabaseError,ValueError) as error: raise ValidationError(f'恢复被拒绝：{error}')
    from django.db.models import F
    m.User.objects.update(auth_version=int(timezone.now().timestamp()))
    audit(None,'恢复数据库','database',details={'safety_backup':recovery})
    return {'message':'已恢复数据，请重新登录。','safety_backup':recovery}
